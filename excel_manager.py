from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import (
    Table,
    TableStyleInfo,
)

from utils import parse_date


# ============================================================
# XLSX 저장
# ============================================================

def save_xlsx(
    clips,
    filename: Path,
):
    """
    클립 데이터를 XLSX로 저장합니다.

    기능:
        - 클릭 가능한 하이퍼링크
        - 실제 Excel 날짜 형식
        - Excel Table
        - 행 줄무늬
        - 필터
        - 첫 행 고정
    """

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "클립 목록"

    # ==============================================
    # 헤더
    # ==============================================

    headers = [
        "제목",
        "링크",
        "클립 UID",
        "날짜",
    ]

    worksheet.append(headers)

    # 헤더 굵게
    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

    # ==============================================
    # 데이터
    # ==============================================

    for clip in clips:

        row_number = (
            worksheet.max_row + 1
        )

        # ------------------------------------------
        # 제목
        # ------------------------------------------

        worksheet.cell(
            row=row_number,
            column=1,
            value=clip["title"],
        )

        # ------------------------------------------
        # 링크
        # ------------------------------------------

        link_cell = worksheet.cell(
            row=row_number,
            column=2,
            value=clip["url"],
        )

        link_cell.hyperlink = clip["url"]
        link_cell.style = "Hyperlink"

        # ------------------------------------------
        # 클립 UID
        # ------------------------------------------

        worksheet.cell(
            row=row_number,
            column=3,
            value=clip["clip_uid"],
        )

        # ------------------------------------------
        # 날짜
        # ------------------------------------------

        date_value = parse_date(
            clip["created_date"]
        )

        date_cell = worksheet.cell(
            row=row_number,
            column=4,
        )

        if date_value is not None:

            # 실제 Excel 날짜
            date_cell.value = date_value

            date_cell.number_format = (
                "yyyy-mm-dd hh:mm:ss"
            )

        else:

            # 날짜 형식이 잘못된 경우
            # 원본 문자열 그대로 저장
            date_cell.value = (
                clip["created_date"]
            )

    # ==============================================
    # Excel Table
    # ==============================================

    if worksheet.max_row >= 2:

        table = Table(
            displayName="ClipTable",
            ref=f"A1:D{worksheet.max_row}",
        )

        table_style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        table.tableStyleInfo = table_style

        worksheet.add_table(table)

    # ==============================================
    # 열 너비
    # ==============================================

    worksheet.column_dimensions["A"].width = 50
    worksheet.column_dimensions["B"].width = 55
    worksheet.column_dimensions["C"].width = 30
    worksheet.column_dimensions["D"].width = 22

    # ==============================================
    # 첫 행 고정
    # ==============================================

    worksheet.freeze_panes = "A2"

    # ==============================================
    # 저장
    # ==============================================

    workbook.save(filename)