import csv
import sys
from pathlib import Path

from openpyxl import load_workbook


# ============================================================
# 프로그램 기본 폴더
# ============================================================

def get_base_directory():
    """
    프로그램이 실행되는 기준 폴더를 반환합니다.
    """

    if getattr(sys, "frozen", False):
        return Path(
            sys.executable
        ).resolve().parent

    return Path(
        __file__
    ).resolve().parent


# ============================================================
# files 폴더
# ============================================================

def get_files_directory():
    """
    exe 또는 Python 파일과 같은 위치의
    files 폴더를 반환합니다.

    없으면 자동 생성합니다.
    """

    base_directory = get_base_directory()

    files_directory = (
        base_directory
        / "files"
    )

    files_directory.mkdir(
        parents=True,
        exist_ok=True,
    )

    return files_directory


# ============================================================
# 기존 파일 찾기
# ============================================================

def find_existing_file(
    channel_code: str,
    files_directory: Path,
):
    """
    files 폴더에서 해당 채널의 기존 파일을 찾습니다.

    우선순위:
        1. XLSX
        2. CSV

    여러 개라면 가장 최근 수정된 파일을 사용합니다.
    """

    xlsx_pattern = (
        f"*_{channel_code}.xlsx"
    )

    csv_pattern = (
        f"*_{channel_code}.csv"
    )

    xlsx_files = list(
        files_directory.glob(
            xlsx_pattern
        )
    )

    if xlsx_files:
        return max(
            xlsx_files,
            key=lambda file: file.stat().st_mtime
        )

    csv_files = list(
        files_directory.glob(
            csv_pattern
        )
    )

    if csv_files:
        return max(
            csv_files,
            key=lambda file: file.stat().st_mtime
        )

    return None


# ============================================================
# 전체 기존 채널 파일 찾기
# ============================================================

def find_all_channel_files(
    files_directory: Path,
):
    """
    files 폴더에 존재하는 모든 채널 파일을 찾습니다.

    파일명 형식:
        채널이름_날짜_시간_채널코드.xlsx

    예:
        침착맨_20261007_010000_ABC123.xlsx

    마지막 '_' 뒤의 값을 채널 코드로 사용합니다.

    우선순위:
        1. XLSX
        2. CSV

    같은 채널 코드의 파일이 여러 개라면
    가장 적절한 파일 하나만 반환합니다.
    """

    channel_files = {}

    for file in files_directory.iterdir():

        if not file.is_file():
            continue

        suffix = file.suffix.lower()

        if suffix not in {
            ".xlsx",
            ".csv",
        }:
            continue

        # 파일명에서 마지막 "_" 뒤의 값을 채널 코드로 사용
        parts = file.stem.rsplit(
            "_",
            1,
        )

        if len(parts) != 2:
            continue

        channel_code = parts[1].strip()

        if not channel_code:
            continue

        current_file = channel_files.get(
            channel_code
        )

        # 처음 발견한 파일
        if current_file is None:

            channel_files[channel_code] = file

            continue

        current_suffix = (
            current_file.suffix.lower()
        )

        # XLSX를 CSV보다 우선
        if (
            suffix == ".xlsx"
            and current_suffix == ".csv"
        ):

            channel_files[channel_code] = file

            continue

        # 같은 확장자라면 최신 파일 우선
        if suffix == current_suffix:

            if file.stat().st_mtime > current_file.stat().st_mtime:

                channel_files[channel_code] = file

    # 파일명 기준으로 정렬
    return sorted(
        channel_files.items(),
        key=lambda item: item[1].name.lower(),
    )


# ============================================================
# 기존 XLSX 읽기
# ============================================================

def load_existing_xlsx(filename: Path):
    """
    기존 XLSX 파일에서 클립 데이터를 읽습니다.
    """

    clips = []

    try:

        workbook = load_workbook(
            filename,
            read_only=True,
            data_only=False,
        )

        worksheet = workbook.active

        # --------------------------------------------
        # 헤더 위치 확인
        # --------------------------------------------

        headers = {}

        for cell in worksheet[1]:

            if cell.value is not None:
                headers[str(cell.value)] = (
                    cell.column
                )

        title_column = headers.get(
            "제목"
        )

        link_column = headers.get(
            "링크"
        )

        uid_column = headers.get(
            "클립 UID"
        )

        date_column = headers.get(
            "날짜"
        )

        if uid_column is None:

            print(
                "XLSX에서 '클립 UID' 열을 "
                "찾지 못했습니다."
            )

            workbook.close()

            return []

        # --------------------------------------------
        # 데이터 읽기
        # --------------------------------------------

        for row in worksheet.iter_rows(
            min_row=2
        ):

            def get_value(column):

                if column is None:
                    return ""

                value = row[
                    column - 1
                ].value

                if value is None:
                    return ""

                return str(value).strip()

            clip_uid = get_value(
                uid_column
            )

            if not clip_uid:
                continue

            created_date = get_value(
                date_column
            )

            # Excel datetime이 문자열로 읽힌 경우
            # 기존 형식으로 통일
            try:

                from datetime import datetime

                parsed_date = datetime.strptime(
                    created_date,
                    "%Y-%m-%d %H:%M:%S"
                )

                created_date = parsed_date.strftime(
                    "%Y-%m-%d %H:%M:%S"
                )

            except ValueError:
                pass

            clips.append({
                "title": get_value(
                    title_column
                ),

                "url": get_value(
                    link_column
                ),

                "clip_uid": clip_uid,

                "created_date": created_date,
            })

        workbook.close()

    except Exception as e:

        print(
            f"XLSX 파일을 읽지 못했습니다: {e}"
        )

        return []

    return clips


# ============================================================
# 기존 CSV 읽기
# ============================================================

def load_existing_csv(filename: Path):
    """
    기존 CSV 파일에서 클립 데이터를 읽습니다.
    """

    clips = []

    try:

        with open(
            filename,
            "r",
            newline="",
            encoding="utf-8-sig",
        ) as file:

            reader = csv.DictReader(
                file
            )

            for row in reader:

                clip_uid = row.get(
                    "클립 UID",
                    ""
                ).strip()

                if not clip_uid:
                    continue

                clips.append({
                    "title": row.get(
                        "제목",
                        ""
                    ).strip(),

                    "url": row.get(
                        "링크",
                        ""
                    ).strip(),

                    "clip_uid": clip_uid,

                    "created_date": row.get(
                        "날짜",
                        ""
                    ).strip(),
                })

    except OSError as e:

        print(
            f"CSV 파일을 읽지 못했습니다: {e}"
        )

        return []

    return clips


# ============================================================
# 기존 파일 데이터 읽기
# ============================================================

def load_existing_clips(filename):
    """
    기존 XLSX 또는 CSV 파일을 읽습니다.
    """

    if filename is None:
        return []

    suffix = filename.suffix.lower()

    if suffix == ".xlsx":

        return load_existing_xlsx(
            filename
        )

    if suffix == ".csv":

        return load_existing_csv(
            filename
        )

    return []